import os
from datetime import datetime

from docx import Document

from core.utils import sanitize_filename
from documents import pdf_generator


def generate_cv_documents(cv_text, company, job_title, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    stamp = datetime.now().strftime('%Y%m%d')
    base = sanitize_filename(f'CV_{company}_{job_title}_{stamp}')

    docx_path = os.path.join(out_dir, f'{base}.docx')
    doc = Document()
    for line in (cv_text or '').splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            doc.add_heading(line[2:], level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)
    doc.save(docx_path)

    pdf_path = os.path.join(out_dir, f'{base}.pdf')
    pdf_generator.html_to_pdf(pdf_generator.markdown_to_html(cv_text), pdf_path)
    return {'docx': docx_path, 'pdf': pdf_path, 'version': base}
